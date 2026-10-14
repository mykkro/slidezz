from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

from backend.templates.base import SlideTemplate
from backend.paths import PROJECT_ROOT


class TitleTemplate(SlideTemplate):
    type_name = "title"

    def render(self, prs, slide_data, theme):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        image_name = slide_data.get(
            "background_image"
        )

        if image_name:
            image_path = (
                PROJECT_ROOT /
                "images" /
                image_name
            )

            if image_path.exists():
                slide.shapes.add_picture(
                    str(image_path),
                    0,
                    0,
                    prs.slide_width,
                    prs.slide_height,
                )

                """
                overlay = slide.shapes.add_shape(
                    MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                    0,
                    0,
                    prs.slide_width,
                    prs.slide_height,
                )

                overlay.fill.solid()
                overlay.fill.fore_color.rgb = RGBColor(
                    0,
                    0,
                    0,
                )

                overlay.fill.transparency = 45

                overlay.line.fill.background()
                """

        title_box = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(1.8),
            Inches(10),
            Inches(1),
        )

        title_frame = title_box.text_frame
        title_frame.clear()

        title_paragraph = title_frame.paragraphs[0]
        title_paragraph.text = slide_data.get(
            "title",
            "",
        )

        title_font = title_paragraph.font
        title_font.name = theme["font_name"]
        title_font.size = Pt(
            theme["title_size"]
        )

        if image_name:
            title_font.color.rgb = RGBColor(
                255,
                255,
                255,
            )
        else:
            title_font.color.rgb = RGBColor.from_string(
                theme["primary_color"].replace(
                    "#",
                    "",
                )
            )

        subtitle_box = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(3.0),
            Inches(10),
            Inches(1),
        )

        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.clear()

        subtitle_paragraph = (
            subtitle_frame.paragraphs[0]
        )

        subtitle_paragraph.text = slide_data.get(
            "subtitle",
            "",
        )

        subtitle_font = subtitle_paragraph.font
        subtitle_font.name = theme["font_name"]
        subtitle_font.size = Pt(
            theme["subtitle_size"]
        )

        if image_name:
            subtitle_font.color.rgb = RGBColor(
                255,
                255,
                255,
            )
        else:
            subtitle_font.color.rgb = RGBColor(
                80,
                80,
                80,
            )

        author = slide_data.get(
            "author",
            "",
        )

        institution = slide_data.get(
            "institution",
            "",
        )

        date = slide_data.get(
            "date",
            "",
        )

        footer_lines = []

        if author:
            footer_lines.append(author)

        if institution:
            footer_lines.append(institution)

        if date:
            footer_lines.append(date)

        footer_text = "\n".join(
            footer_lines
        )

        if footer_text:
            footer_box = slide.shapes.add_textbox(
                Inches(0.8),
                Inches(4.8),
                Inches(6),
                Inches(1.2),
            )

            footer_frame = footer_box.text_frame
            footer_frame.clear()

            footer_frame.paragraphs[0].text = footer_text

            footer_font = (
                footer_frame.paragraphs[0].font
            )

            footer_font.name = theme["font_name"]
            footer_font.size = Pt(16)

            if image_name:
                footer_font.color.rgb = RGBColor(
                    255,
                    255,
                    255,
                )
            else:
                footer_font.color.rgb = RGBColor(
                    100,
                    100,
                    100,
                )