from pptx.util import Inches

from backend.templates.base import SlideTemplate


class ContentLeftImageTemplate(SlideTemplate):
    type_name = "content_left_image"

    supports_background = True

    schema = [
        {
            "name": "title",
            "type": "text",
            "label": "Title",
        },
        {
            "name": "image",
            "type": "image",
            "label": "Image",
        },
        {
            "name": "markdown",
            "type": "markdown",
            "label": "Content",
        },
    ]

    def render(
        self,
        prs,
        slide_data,
        theme,
        page_number,
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        self.add_background(
            prs,
            slide,
            slide_data,
        )

        self.add_title(
            slide,
            slide_data.get(
                "title",
                ""
            ),
            theme,
        )

        image_path = self.get_image_path(
            slide_data.get("image")
        )

        slide.shapes.add_picture(
            str(image_path),
            Inches(7.0),
            Inches(1.3),
            width=Inches(5),
        )

        content_box = (
            slide.shapes.add_textbox(
                Inches(0.5),
                Inches(1.3),
                Inches(6.5),
                Inches(4.5),
            )
        )

        self.fill_markdown(
            content_box,
            slide_data["markdown"]
        )
