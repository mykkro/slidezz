from pptx.util import Inches
from pptx.util import Pt

from backend.templates.base import SlideTemplate

class ContentTemplate(SlideTemplate):

    type_name = "content"

    supports_background = True

    schema = [
        {
            "name": "title",
            "type": "text",
            "label": "Title",
        },
        {
            "name": "markdown",
            "type": "markdown",
            "label": "Content",
        },
        {
            "name": "background_image",
            "type": "image",
            "label": "Background Image",
        },
    ]

    def render(
        self,
        prs,
        slide_data,
        theme,
        page_number,
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        if self.supports_background:
            self.add_background(
                prs=prs,
                slide=slide,
                slide_data=slide_data,
            )

        title_box = slide.shapes.add_textbox(
            Inches(0.5),
            Inches(0.3),
            Inches(11),
            Inches(0.5),
        )

        title_box.text_frame.text = slide_data["title"]

        content_box = slide.shapes.add_textbox(
            Inches(0.7),
            Inches(1.2),
            Inches(11),
            Inches(4.5),
        )

        content_box.text_frame.text = slide_data.get(
            "markdown",
            "",
        )

        self.add_footer(
            slide,
            prs,
            theme.get(
                "footer_text",
                "",
            ),
        )

        self.add_page_number(
            slide,
            prs,
            page_number,
        )