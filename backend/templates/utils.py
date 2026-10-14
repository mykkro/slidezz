from pathlib import Path

from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE


def add_background_picture(
    slide,
    image_path,
):
    slide.shapes.add_picture(
        str(image_path),
        0,
        0,
        width=slide.part.presentation.slide_width,
        height=slide.part.presentation.slide_height,
    )