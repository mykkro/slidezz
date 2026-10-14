"""Native editable PPTX shapes, using the same canvas as the HTML preview."""
import html
from io import BytesIO
import re

import markdown
from bs4 import BeautifulSoup, NavigableString
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from backend.paths import PROJECT_ROOT
from backend.templates.registry import IMAGE_ALIGNMENTS
from backend.services.layout_geometry import DEFAULTS as LAYOUT_DEFAULTS, resolve_layout

DEFAULT_THEME = dict(font_name="Aptos", primary_color="#16745C", background_color="#FFFFFF",
                     title_size=30, subtitle_size=18, footer_text="", text_color="#263631",
                     footer_color="#263631", background_panel_color="#FFFFFF",
                     background_panel_transparency=50)
DEFAULT_THEME.update(LAYOUT_DEFAULTS)


def image_path(name, images_dir=None):
    root = (images_dir or (PROJECT_ROOT / "images")).resolve()
    path = (root / name).resolve()
    return path if path.parent == root and path.is_file() else None


def markdown_html(text):
    rendered = markdown.markdown(html.escape(text), extensions=["fenced_code", "sane_lists"])
    soup = BeautifulSoup(rendered, "html.parser")
    for link in soup.find_all("a"):
        if not re.match(r"^(https?://|mailto:|#)", link.get("href", ""), re.I):
            link.attrs.pop("href", None)
    for image in soup.find_all("img"):
        image.replace_with(f"[Image {image.get('alt', '')}: use a gallery image field]")
    return str(soup)


def picture_source(path):
    if path.suffix.lower() == ".webp":
        stream = BytesIO()
        with Image.open(path) as image:
            image.save(stream, format="PNG")
        stream.seek(0)
        return stream
    return str(path)


def render_pptx(prs, template, data, theme, number, images_dir=None):
    def resolve_image(name):
        return image_path(name, images_dir)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(theme["background_color"][1:])
    background = data.get("background_image", "")
    geometry = resolve_layout(template, theme)
    if background and resolve_image(background):
        slide.shapes.add_picture(picture_source(resolve_image(background)), 0, 0, prs.slide_width, prs.slide_height)
        surface = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                        *(Inches(geometry["surface"][key]) for key in ("x", "y", "w", "h")))
        surface.fill.solid()
        surface.fill.fore_color.rgb = RGBColor.from_string(data.get("background_panel_color", theme["background_panel_color"])[1:])
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(round((100 - data.get("background_panel_transparency", theme["background_panel_transparency"])) * 1000)))
        surface._element.spPr.xpath("./a:solidFill/a:srgbClr")[0].append(alpha)
        surface.line.fill.background()
    for item in geometry["layout"]:
        value = data.get(item["field"], "")
        x, y, w, h = (Inches(item[key]) for key in ("x", "y", "w", "h"))
        if item["kind"] == "image" and value and resolve_image(value):
            fit = data.get(item["field"] + "_fit", "contain")
            alignment = next(a for a in IMAGE_ALIGNMENTS
                             if a["value"] == data.get(item["field"] + "_alignment", "center"))
            ax, ay = alignment["x"], alignment["y"]
            with Image.open(resolve_image(value)) as image:
                ratio = (max if fit == "cover" else min)(w / image.width, h / image.height)
                iw, ih = image.width * ratio, image.height * ratio
            if fit == "cover":
                picture = slide.shapes.add_picture(picture_source(resolve_image(value)), x, y, w, h)
                crop_x, crop_y = max(0, 1 - w / iw), max(0, 1 - h / ih)
                picture.crop_left, picture.crop_right = crop_x * ax, crop_x * (1 - ax)
                picture.crop_top, picture.crop_bottom = crop_y * ay, crop_y * (1 - ay)
            else:
                slide.shapes.add_picture(picture_source(resolve_image(value)),
                                        x + round((w-iw)*ax), y + round((h-ih)*ay),
                                        round(iw), round(ih))
            continue
        if item["kind"] == "image":
            value = f"Image {value or '(not selected)'} not found"
        frame = slide.shapes.add_textbox(x, y, w, h).text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        markup = markdown_html(value) if item["kind"] == "markdown" else f"<p>{html.escape(value)}</p>"
        soup = BeautifulSoup(markup, "html.parser")
        # List items own their paragraphs; nested items are rendered separately.
        blocks = [b for b in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre"])
                  if b.name == "li" or not b.find_parent("li")]
        for i, block in enumerate(blocks):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.space_after = Pt(10)
            heading = block.name.startswith("h")
            paragraph.font.name = theme["font_name"]
            paragraph.font.size = Pt(item["size"] + (5 if heading else 0))
            paragraph.font.color.rgb = RGBColor.from_string(
                (theme["primary_color"] if item["field"] == "title" or heading else theme["text_color"])[1:])
            if block.name == "li":
                parent = block.parent
                depth = len([p for p in block.parents if p.name in ("ul", "ol")])
                paragraph._p.get_or_add_pPr().set("marL", str(Inches(.3125 * depth)))
                paragraph._p.get_or_add_pPr().set("indent", str(-Inches(.15)))
                start = int(parent.get("start", 1))
                prefix = f"{list(parent.find_all('li', recursive=False)).index(block) + start}. " if parent.name == "ol" else "\u2022 "
                paragraph.add_run().text = prefix
            for descendant in block.descendants:
                if isinstance(descendant, NavigableString):
                    if block.name == "li" and next((p for p in descendant.parents if p.name == "li"), None) is not block:
                        continue
                    run = paragraph.add_run()
                    run.text = str(descendant)
                    parents = list(descendant.parents)
                    run.font.bold = heading or any(p.name in ("strong", "b") for p in parents)
                    run.font.italic = any(p.name in ("em", "i") for p in parents)
                    if any(p.name in ("code", "pre") for p in parents):
                        run.font.name = "Consolas"
                    link = next((p for p in parents if p.name == "a"), None)
                    if link and link.get("href"):
                        run.hyperlink.address = link["href"]
                elif descendant.name == "br":
                    paragraph.add_line_break()
    footer = geometry["footer"]
    for text, x, width in [(theme.get("footer_text", ""), footer["x"], footer["w"]),
                           (f"{number:02d}", footer["page_x"], .65)]:
        frame = slide.shapes.add_textbox(Inches(x), Inches(footer["y"]), Inches(width), Inches(.3)).text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.text = text
        frame.paragraphs[0].font.name = theme["font_name"]
        frame.paragraphs[0].font.size = Pt(10)
        frame.paragraphs[0].font.color.rgb = RGBColor.from_string(theme["footer_color"][1:])
    if background and not resolve_image(background):
        frame = slide.shapes.add_textbox(Inches(footer["x"]), Inches(geometry["missing_y"]), Inches(footer["w"]), Inches(.3)).text_frame
        frame.text = f"Image {background} not found"
        frame.paragraphs[0].font.size = Pt(10)
