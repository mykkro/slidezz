"""Validate the published YAML examples and coverage of the editable schema."""
from io import BytesIO
from pathlib import Path
import re
import yaml
from pptx import Presentation
from backend.app import create_app, validate
from backend.services.pptx_generator import generate
from backend.services.slide_renderer import DEFAULT_THEME
from backend.templates.registry import get_template_definitions

SOURCE = (Path(__file__).resolve().parents[1] / "AUTHORING.md").read_text(encoding="utf-8")
EXAMPLES = re.findall(r"```yaml\n(.*?)\n```", SOURCE, re.S)


def test_guide_yaml_examples_import_and_export(tmp_path):
    app = create_app({"TESTING": True, "PRESENTATION_FILE": tmp_path / "deck.json",
                      "IMAGES_DIR": tmp_path / "images"})
    client = app.test_client()
    assert len(EXAMPLES) == 4
    for source in EXAMPLES:
        response = client.post("/api/import/yaml", json={"text": source})
        assert response.status_code == 200, response.json
        deck = validate(response.json)
        output = BytesIO()
        generate(deck, output)
        output.seek(0)
        assert len(Presentation(output).slides) == len(deck["slides"])


def test_programmatic_authoring_examples(tmp_path, monkeypatch):
    import json

    monkeypatch.chdir(tmp_path)
    snippets = re.findall(r"```python\n(.*?)\n```", SOURCE, re.S)
    assert len(snippets) == 2
    for index, snippet in enumerate(snippets):
        exec(compile(snippet, f"AUTHORING.md:python-example-{index + 1}", "exec"), {})
    generated = yaml.safe_load((tmp_path / "generated.yaml").read_text(encoding="utf-8"))
    assert generated == json.loads((tmp_path / "generated.json").read_text(encoding="utf-8"))
    assert len(generated["slides"]) == 7
    app = create_app({"TESTING": True, "PRESENTATION_FILE": tmp_path / "saved.json",
                      "IMAGES_DIR": tmp_path / "images"})
    response = app.test_client().post("/api/import/yaml", json={
        "text": (tmp_path / "generated.yaml").read_text(encoding="utf-8")})
    assert response.status_code == 200, response.json
    assert len(Presentation(tmp_path / "generated.pptx").slides) == 7


def test_guide_covers_all_templates_and_fields():
    example = yaml.safe_load(EXAMPLES[0])
    definitions = get_template_definitions()
    assert {slide["type"] for slide in example["slides"]} == set(definitions)
    assert set(example["theme"]) == set(DEFAULT_THEME)
    for name, definition in definitions.items():
        assert f"`{name}`" in SOURCE
        for field in definition["fields"]:
            key = field["name"]
            if key.endswith("_vertical_alignment"):
                assert "`<field>_vertical_alignment`" in SOURCE
            elif key.endswith("_alignment") and field.get("theme_key"):
                assert "`<field>_alignment`" in SOURCE
            elif "_font_" in key:
                base, suffix = key.rsplit("_font_", 1)
                assert f"`{base}`" in SOURCE
                assert f"`<field>_font_{suffix}`" in SOURCE
            else:
                assert f"`{key}`" in SOURCE, (name, key)
