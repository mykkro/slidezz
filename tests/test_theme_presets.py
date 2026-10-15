from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from bs4 import BeautifulSoup

from backend.app import create_app, validate
from backend.services.bundles import export_zip, referenced_images, build_static_html
from backend.services.pptx_generator import generate
from backend.templates.registry import get_all_templates

ROOT = Path(__file__).resolve().parents[1]
PRESETS = json.loads((ROOT / "frontend/assets/themes.json").read_text())


def luminance(color):
    channels = [int(color[i:i+2], 16) / 255 for i in (1, 3, 5)]
    linear = [value / 12.92 if value <= .04045 else ((value + .055) / 1.055) ** 2.4 for value in channels]
    return sum(value * weight for value, weight in zip(linear, (.2126, .7152, .0722)))


def test_presets_validate_have_assets_and_render_all_layouts():
    assert len(PRESETS) == 6
    assert len({preset["id"] for preset in PRESETS}) == 6
    for preset in PRESETS:
        theme = preset["theme"]
        assert (ROOT / "images" / theme["default_background_image"]).is_file()
        for key in ("text_color", "primary_color"):
            values = sorted([luminance(theme[key]), luminance(theme["background_panel_color"])])
            assert (values[1] + .05) / (values[0] + .05) >= 4.5
        slides = [{"type":name, **{field["name"]:field["default"] for field in template.schema if not field.get("theme_key")},
                   "background_image":theme["default_background_image"]} for name, template in get_all_templates().items()]
        output = BytesIO()
        generate(validate({"theme":theme, "slides":slides}), output, ROOT / "images")
        output.seek(0)
        deck = Presentation(output)
        assert len(deck.slides) == len(get_all_templates())
        assert all(len(slide.shapes) >= 4 for slide in deck.slides)


def test_default_background_in_bundle_and_collision_remapping(tmp_path):
    raw = BytesIO()
    Image.new("RGB", (10, 10), "red").save(raw, "PNG")
    (tmp_path / "custom.png").write_bytes(raw.getvalue())
    data = validate({"theme":{"default_background_image":"custom.png"}, "slides":[]})
    assert referenced_images(data) == ["custom.png"]
    bundle = export_zip(data, tmp_path)
    with ZipFile(BytesIO(bundle)) as archive:
        assert archive.read("images/custom.png") == raw.getvalue()
    target = tmp_path / "target"
    target.mkdir()
    (target / "custom.png").write_bytes(b"different")
    app = create_app({"TESTING":True,"IMAGES_DIR":target,"PRESENTATION_FILE":tmp_path / "deck.json"})
    result = app.test_client().post("/api/import/zip", data={"archive":(BytesIO(bundle), "deck.zip")})
    assert result.status_code == 200
    renamed = result.json["theme"]["default_background_image"]
    assert renamed != "custom.png"
    assert (target / renamed).read_bytes() == raw.getvalue()


def test_standalone_contains_all_preset_assets_and_rejects_invalid_default(tmp_path):
    soup = BeautifulSoup(build_static_html(validate({"slides":[]}), {}), "html.parser")
    config = json.loads(soup.select_one("#static-config").string)
    assert config["presets"] == PRESETS
    for preset in PRESETS:
        assert config["images"][preset["theme"]["default_background_image"]].startswith("data:image/png;base64,")
    app = create_app({"TESTING":True,"IMAGES_DIR":tmp_path,"PRESENTATION_FILE":tmp_path / "deck.json"})
    assert app.test_client().get("/api/theme/presets").json == PRESETS
    for value in (None, [], 5, "x" * 201):
        assert app.test_client().post("/api/save", json={"theme":{"default_background_image":value},"slides":[]}).status_code == 400
