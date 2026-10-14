from io import BytesIO
from pathlib import Path
import pytest
from PIL import Image
from pptx import Presentation

from backend.app import create_app, validate
from backend.templates.registry import get_all_templates, IMAGE_ALIGNMENTS
from backend.services.slide_renderer import markdown_html
from backend.services.pptx_generator import generate


@pytest.fixture
def client(tmp_path):
    app = create_app({"TESTING": True, "PRESENTATION_FILE": tmp_path / "deck.json",
                      "IMAGES_DIR": tmp_path / "images"})
    return app.test_client()


def test_all_schemas_export_valid_widescreen_shapes():
    slides = []
    for name, template in get_all_templates().items():
        assert {b["field"] for b in template.layout} <= {f["name"] for f in template.schema}
        slides.append({"type": name, **{f["name"]: f["default"] for f in template.schema}})
    out = BytesIO()
    generate(validate({"slides": slides}), out)
    out.seek(0)
    prs = Presentation(out)
    assert len(prs.slides) == 8
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < .001
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height


def test_markdown_is_rich_and_safe():
    html = markdown_html("## Heading\n\n**Bold** and *italic*\n\n- one\n- two\n\n<script>alert(1)</script>\n\n[x](javascript:alert)")
    assert "<h2>" in html and "<strong>" in html and "<em>" in html and "<li>" in html
    assert "<script>" not in html and "href=\"javascript:" not in html


def test_pptx_contains_formatted_runs_and_missing_image_message(client):
    data = {"slides": [
        {"type": "content", "title": "Review", "markdown": "**Bold** and *italic*\n\n1. First\n2. Second",
         "background_image": "missing-bg.png"},
        {"type": "double_image_content", "left_image": "missing.png", "right_image": "../../README.md"},
        {"type": "title", "title": "Title", "author": "Author", "institution": "Org"}
    ]}
    result = client.post("/api/export/pptx", json=data)
    assert result.status_code == 200
    prs = Presentation(BytesIO(result.data))
    text = "\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert "Image missing.png not found" in text
    assert "Image missing-bg.png not found" in text
    assert "1. First" in text and "2. Second" in text
    runs = [r for s in prs.slides[0].shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs]
    assert any(r.text == "Bold" and r.font.bold for r in runs)
    assert any(r.text == "italic" and r.font.italic for r in runs)


@pytest.mark.parametrize("data", [
    {}, {"slides": {}}, {"slides": [{"type": "unknown"}]},
    {"slides": [{"type": "content", "markdown": 3}]},
    {"slides": [], "theme": {"primary_color": "red"}},
    {"slides": [], "theme": []},
    {"slides": [{"type": []}]},
    {"slides": [{"type": "title"}] * 201},
])
def test_invalid_import_does_not_replace_saved_deck(client, data):
    original = {"slides": [{"type": "title", "title": "Keep me"}]}
    assert client.post("/api/save", json=original).status_code == 200
    response = client.post("/api/import", json=data)
    assert response.status_code == 400 and response.json["error"]
    assert client.get("/api/presentation").json["slides"][0]["title"] == "Keep me"


def test_import_and_save_round_trip(client):
    imported = client.post("/api/import", json={"slides": [{"type": "two_columns", "left_markdown": "**Left**"}]})
    assert imported.status_code == 200
    assert imported.json["slides"][0]["right_markdown"] == ""
    assert client.post("/api/save", json=imported.json).status_code == 200
    assert client.get("/api/presentation").json == imported.json


def test_upload_validation_unique_names_and_gallery(client):
    raw = BytesIO()
    Image.new("RGB", (120, 80), "green").save(raw, "PNG")
    names = []
    for _ in range(2):
        response = client.post("/api/images", data={"image": (BytesIO(raw.getvalue()), "../../same.png")})
        assert response.status_code == 201
        names.append(response.json["name"])
    assert names[0] != names[1]
    assert "/" not in names[0] and "\\" not in names[0]
    assert client.get("/api/images").json == sorted(names)
    assert client.get("/images/" + names[0]).status_code == 200
    assert client.post("/api/images", data={"image": (BytesIO(b"invalid"), "fake.png")}).status_code == 400
    assert client.get("/images/../../README.md").status_code == 404


def test_local_assets_and_malformed_requests(client):
    html = client.get("/").text
    assert "cdn.jsdelivr" not in html and "https://" not in html
    for asset in ("bootstrap/bootstrap.min.css", "bootstrap/bootstrap.bundle.min.js",
                  "easymde/easymde.min.js", "lucide/lucide.min.js"):
        assert client.get("/static/vendor/" + asset).status_code == 200
    assert client.post("/api/save", data="{", content_type="application/json").status_code == 400
    assert client.post("/api/markdown", json={"text": 3}).status_code == 400


def test_uploaded_webp_exports_from_configured_gallery(client):
    raw = BytesIO()
    Image.new("RGB", (120, 80), "green").save(raw, "WEBP")
    uploaded = client.post("/api/images", data={"image": (BytesIO(raw.getvalue()), "photo.webp")}).json["name"]
    response = client.post("/api/export/pptx", json={"slides": [
        {"type": "image_left_content", "image": uploaded},
        {"type": "title", "title": "With background", "background_image": uploaded},
    ]})
    assert response.status_code == 200
    prs = Presentation(BytesIO(response.data))
    for slide in prs.slides:
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        assert len(pictures) == 1
        assert pictures[0].image.content_type == "image/png"


def test_nested_markdown_lists_preserve_parent_and_child():
    out = BytesIO()
    generate(validate({"slides": [{"type": "content", "markdown":
        "- Parent\n    - Child\n- Next\n\n3. Third\n4. Fourth"}]}), out)
    out.seek(0)
    prs = Presentation(out)
    text = "\n".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert text.count("Parent") == 1 and text.count("Child") == 1 and text.count("Next") == 1
    assert "3. Third" in text and "4. Fourth" in text


def test_existing_slide_metadata_and_background_survive_save(client):
    data = {"slides": [{"type": "image_left_content", "title": "Existing",
                       "background_image": "smoke-blue-blush.png", "subtitle": "Legacy metadata"}]}
    response = client.post("/api/save", json=data)
    assert response.status_code == 200
    slide = client.get("/api/presentation").json["slides"][0]
    assert slide["subtitle"] == "Legacy metadata"
    assert slide["background_image"] == "smoke-blue-blush.png"


@pytest.mark.parametrize("dimensions", [(200, 100), (100, 200)])
@pytest.mark.parametrize("fit", ["contain", "cover"])
@pytest.mark.parametrize("alignment", IMAGE_ALIGNMENTS, ids=lambda a: a["value"])
def test_image_pane_geometry_and_crop(tmp_path, dimensions, fit, alignment):
    Image.new("RGB", dimensions, "green").save(tmp_path / "pane.png")
    data = validate({"slides": [{"type": "image_left_content", "image": "pane.png",
                                "image_fit": fit, "image_alignment": alignment["value"]}]})
    output = BytesIO()
    generate(data, output, tmp_path)
    output.seek(0)
    picture = next(shape for shape in Presentation(output).slides[0].shapes if shape.shape_type == 13)
    from pptx.util import Inches
    x, y, w, h = (Inches(value) for value in (.65, 1.65, 5.7, 4.9))
    if fit == "contain":
        assert picture.crop_left == picture.crop_right == picture.crop_top == picture.crop_bottom == 0
        assert picture.width / picture.height == pytest.approx(dimensions[0] / dimensions[1], abs=1e-6)
        assert picture.left == pytest.approx(x + (w - picture.width) * alignment["x"], abs=1)
        assert picture.top == pytest.approx(y + (h - picture.height) * alignment["y"], abs=1)
        assert picture.width <= w and picture.height <= h
    else:
        assert (picture.left, picture.top, picture.width, picture.height) == (x, y, w, h)
        crop_x = picture.crop_left + picture.crop_right
        crop_y = picture.crop_top + picture.crop_bottom
        assert picture.crop_left == pytest.approx(crop_x * alignment["x"], abs=1e-5)
        assert picture.crop_top == pytest.approx(crop_y * alignment["y"], abs=1e-5)
        uncropped_ratio = (w / (1 - crop_x)) / (h / (1 - crop_y))
        assert uncropped_ratio == pytest.approx(dimensions[0] / dimensions[1], abs=1e-4)


def test_image_settings_defaults_round_trip_and_validation(client):
    imported = client.post("/api/import", json={"slides": [{"type": "double_image_content"}]}).json
    slide = imported["slides"][0]
    assert slide["left_image_fit"] == slide["right_image_fit"] == "contain"
    assert slide["left_image_alignment"] == slide["right_image_alignment"] == "center"
    slide.update(left_image_fit="cover", left_image_alignment="bottom_left",
                 right_image_fit="contain", right_image_alignment="top_right")
    assert client.post("/api/save", json=imported).status_code == 200
    assert client.get("/api/presentation").json == imported
    for key, bad in [("left_image_fit", "stretch"), ("right_image_alignment", "invalid")]:
        invalid = {"slides": [dict(slide, **{key: bad})]}
        response = client.post("/api/import", json=invalid)
        assert response.status_code == 400 and key in response.json["error"]


@pytest.mark.parametrize("transparency", [0, 50, 75, 100])
def test_background_panel_color_and_native_transparency(tmp_path, transparency):
    Image.new("RGB", (200, 100), "blue").save(tmp_path / "background.png")
    data = validate({"slides": [{"type": "title", "background_image": "background.png",
                                "background_panel_color": "#DAF3E6",
                                "background_panel_transparency": transparency}]})
    output = BytesIO()
    generate(data, output, tmp_path)
    output.seek(0)
    slide = Presentation(output).slides[0]
    panel = slide.shapes[1]
    assert str(panel.fill.fore_color.rgb) == "DAF3E6"
    alpha = panel._element.spPr.xpath("./a:solidFill/a:srgbClr/a:alpha")[0]
    assert int(alpha.get("val")) == (100 - transparency) * 1000
    assert not any(shape._element.xpath(".//a:alpha") for shape in list(slide.shapes)[2:])


def test_panel_defaults_and_round_trip(client):
    data = client.post("/api/import", json={"slides": [{"type": "title"}]}).json
    slide = data["slides"][0]
    assert "background_panel_color" not in slide
    assert "background_panel_transparency" not in slide
    assert data["theme"]["background_panel_color"] == "#FFFFFF"
    assert data["theme"]["background_panel_transparency"] == 50
    slide.update(background_panel_color="#ABCDEF", background_panel_transparency=0)
    assert client.post("/api/save", json=data).status_code == 200
    assert client.get("/api/presentation").json == data


@pytest.mark.parametrize("field,value", [
    ("background_panel_color", "red"), ("background_panel_color", "#123"),
    ("background_panel_transparency", -1), ("background_panel_transparency", 101),
    ("background_panel_transparency", "50"), ("background_panel_transparency", True),
    ("background_panel_transparency", None),
])
def test_invalid_panel_settings_rejected(client, field, value):
    response = client.post("/api/import", json={"slides": [{"type": "title", field: value}]})
    assert response.status_code == 400 and field in response.json["error"]


def test_layout_icons_match_registered_panel_geometry(client):
    from xml.etree.ElementTree import fromstring
    namespace = {"svg": "http://www.w3.org/2000/svg"}
    definitions = client.get("/api/templates").json
    icons = set()
    for name, template in get_all_templates().items():
        response = client.get(definitions[name]["icon_url"])
        assert response.status_code == 200 and response.mimetype == "image/svg+xml"
        svg = fromstring(response.data)
        assert svg.attrib["viewBox"] == "0 0 160 90"
        assert svg.find("svg:title", namespace).text == template.label + " layout"
        panels = svg.findall("svg:g", namespace)
        assert [g.attrib["data-field"] for g in panels] == [b["field"] for b in template.layout]
        icons.add(response.data)
    assert len(icons) == len(definitions)
    assert client.get("/api/templates/missing/icon.svg").status_code == 404


def test_theme_panel_inheritance_and_slide_override_export(tmp_path):
    Image.new("RGB", (200, 100), "blue").save(tmp_path / "background.png")
    data = validate({"theme": {"background_panel_color": "#DDEEFF",
                              "background_panel_transparency": 70, "text_color": "#112233",
                              "footer_color": "#445566"},
                     "slides": [
                         {"type": "content", "background_image": "background.png", "markdown": "Body"},
                         {"type": "content", "background_image": "background.png", "markdown": "Override",
                          "background_panel_color": "#FFEEDD", "background_panel_transparency": 0},
                     ]})
    assert "background_panel_color" not in data["slides"][0]
    output = BytesIO()
    generate(data, output, tmp_path)
    output.seek(0)
    prs = Presentation(output)
    for index, (color, alpha_value) in enumerate([("DDEEFF", 30000), ("FFEEDD", 100000)]):
        slide = prs.slides[index]
        panel = slide.shapes[1]
        assert str(panel.fill.fore_color.rgb) == color
        assert panel._element.spPr.xpath("./a:solidFill/a:srgbClr/a:alpha")[0].get("val") == str(alpha_value)
        body = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text in ("Body", "Override"))
        assert str(body.text_frame.paragraphs[0].font.color.rgb) == "112233"
        assert str(slide.shapes[-1].text_frame.paragraphs[0].font.color.rgb) == "445566"
    data["theme"]["background_panel_color"] = "#ABCDEF"
    assert "background_panel_color" not in validate(data)["slides"][0]
    assert validate(data)["slides"][1]["background_panel_color"] == "#FFEEDD"


@pytest.mark.parametrize("key,value", [
    ("background_panel_color", "bad"), ("text_color", "#123"),
    ("footer_color", None), ("background_panel_transparency", 101),
    ("background_panel_transparency", True), ("background_panel_transparency", "50"),
])
def test_invalid_theme_appearance_rejected(client, key, value):
    response = client.post("/api/import", json={"theme": {key: value}, "slides": []})
    assert response.status_code == 400 and key in response.json["error"]


def test_bundled_fonts_and_download_archive(client):
    import json
    from zipfile import ZipFile
    from backend.paths import FRONTEND_DIR
    manifest = json.loads((FRONTEND_DIR / "fonts" / "manifest.json").read_text())
    assert len(manifest["families"]) == 6
    for family in manifest["families"]:
        assert len(family["files"]) == 4
        for name in family["files"]:
            response = client.get(f"/static/fonts/{family['directory']}/{name}")
            assert response.status_code == 200 and response.data[:4] == b"\x00\x01\x00\x00"
        assert client.get(f"/static/fonts/{family['directory']}/OFL.txt").status_code == 200
    download = client.get("/api/fonts/download")
    assert download.status_code == 200
    with ZipFile(BytesIO(download.data)) as archive:
        assert len([name for name in archive.namelist() if name.endswith(".ttf")]) == 24
        assert len([name for name in archive.namelist() if name.endswith("OFL.txt")]) == 6


@pytest.mark.parametrize("font", ["Lato", "Fira Sans", "PT Sans", "PT Serif", "Barlow", "IBM Plex Mono"])
def test_free_font_family_written_to_pptx(font):
    output = BytesIO()
    generate(validate({"theme": {"font_name": font}, "slides": [{"type": "content", "markdown": "**Bold** and *italic*"}]}), output)
    output.seek(0)
    slide = Presentation(output).slides[0]
    for shape in slide.shapes:
        if shape.has_text_frame:
            assert all(p.font.name == font for p in shape.text_frame.paragraphs)


def test_yaml_round_trip_preserves_markdown_theme_and_overrides(client):
    import yaml
    source = 'theme:\n  font_name: Lato\n  background_panel_transparency: 65\nslides:\n  - type: content\n    title: YAML deck\n    background_panel_color: "#abcdef"\n    markdown: |\n      ## Heading\n\n      - **Bold**\n      - Second\n'
    imported = client.post("/api/import/yaml", json={"text": source})
    assert imported.status_code == 200
    exported = client.post("/api/export/yaml", json=imported.json)
    assert exported.status_code == 200 and exported.mimetype == "application/yaml"
    assert "markdown: |" in exported.text
    assert yaml.safe_load(exported.text) == imported.json
    again = client.post("/api/import/yaml", json={"text": exported.text})
    assert again.json == imported.json
    assert client.post("/api/save", json=again.json).status_code == 200
    assert client.get("/api/presentation").json == imported.json


@pytest.mark.parametrize("source", [
    "slides: [", "slides:\n - type: unknown", "slides: nope",
    "!!python/object/apply:os.system ['echo unsafe']",
    "slides: []\nextra: &cycle [*cycle]",
    "slides: []\ntheme:\n  font_name: 2026-01-01",
    "slides: []\nextra: .nan", "slides: []\n1: value",
])
def test_invalid_yaml_does_not_replace_deck(client, source):
    client.post("/api/save", json={"slides": [{"type": "title", "title": "Keep"}]})
    result = client.post("/api/import/yaml", json={"text": source})
    assert result.status_code == 400 and result.json["error"]
    assert client.get("/api/presentation").json["slides"][0]["title"] == "Keep"


def test_theme_geometry_shared_by_preview_and_pptx(client):
    from pptx.util import Inches
    theme = {"margin_left": .9, "margin_right": 1, "margin_top": .7, "margin_bottom": 1.2,
             "column_gap": 1.1, "panel_padding": .3, "footer_margin_bottom": .15,
             "title_size": 36, "body_size": 24}
    layout = client.post("/api/layouts", json=theme).json["two_columns"]
    left, right = layout["layout"][1:3]
    assert left["x"] == .9
    assert right["x"] - left["x"] - left["w"] == pytest.approx(1.1, abs=1e-6)
    assert left["w"] == right["w"]
    response = client.post("/api/export/pptx", json={"theme": theme, "slides": [
        {"type": "two_columns", "title": "Geometry", "left_markdown": "Left", "right_markdown": "Right"}]})
    slide = Presentation(BytesIO(response.data)).slides[0]
    for panel, shape in zip(layout["layout"], list(slide.shapes)[:3]):
        assert tuple(getattr(shape, key) for key in ("left", "top", "width", "height")) == tuple(
            Inches(panel[key]) for key in ("x", "y", "w", "h"))
        assert shape.text_frame.paragraphs[0].font.size.pt == panel["size"]
    assert slide.shapes[-1].top == Inches(layout["footer"]["y"])


@pytest.mark.parametrize("key,value", [
    ("column_gap", -1), ("column_gap", 3), ("margin_left", 4),
    ("margin_bottom", .1), ("panel_padding", True), ("body_size", "24"),
    ("footer_margin_bottom", None), ("title_size", float("nan")),
])
def test_invalid_theme_geometry_rejected(client, key, value):
    result = client.post("/api/import", json={"theme": {key: value}, "slides": []})
    assert result.status_code == 400 and key in result.json["error"]
@pytest.mark.parametrize("original,replacement", [
    ("blue.jpg", "blue-light-mesh.png"),
    ("smoke.jpg", "smoke-blue-blush.png"),
])
def test_retired_background_names_are_migrated(original, replacement):
    data = validate({"slides": [{"type": "image_left_content", "image": original,
                                "background_image": original}]})
    assert data["slides"][0]["image"] == replacement
    assert data["slides"][0]["background_image"] == replacement


def test_manual_covers_templates_and_resolves_curated_images():
    import json
    from backend.services.bundles import referenced_images
    root = Path(__file__).resolve().parents[1]
    data = validate(json.loads((root / "examples" / "manual.json").read_text(encoding="utf-8")))
    assert {slide["type"] for slide in data["slides"]} == set(get_all_templates())
    missing = [name for name in referenced_images(data) if not (root / "images" / name).is_file()]
    assert missing == ["manual-missing-example.png"]
def test_typography_overrides_are_optional_and_exported():
    data = validate({"theme": {"font_name": "Arial"}, "slides": [
        {"type": "content", "title": "Specific title", "title_font_size": 42,
         "title_font_name": "Lato", "title_font_color": "#123456",
         "markdown": "## Heading\n\n**Bold** and *italic*",
         "markdown_font_size": 16, "markdown_font_name": "PT Serif",
         "markdown_font_color": "#654321"},
        {"type": "title", "title": "Inherited title"}]})
    assert "title_font_size" not in data["slides"][1]
    output = BytesIO()
    generate(data, output)
    output.seek(0)
    deck = Presentation(output)
    title = next(shape for shape in deck.slides[0].shapes if shape.has_text_frame and shape.text == "Specific title")
    paragraph = title.text_frame.paragraphs[0]
    assert paragraph.font.size.pt == 42
    assert paragraph.font.name == "Lato"
    assert str(paragraph.font.color.rgb) == "123456"
    body = next(shape for shape in deck.slides[0].shapes if shape.has_text_frame and "Heading" in shape.text)
    assert body.text_frame.paragraphs[0].font.size.pt == 21
    assert body.text_frame.paragraphs[1].font.size.pt == 16
    assert body.text_frame.paragraphs[1].font.name == "PT Serif"
    assert str(body.text_frame.paragraphs[0].font.color.rgb) == "654321"
    inherited = next(shape for shape in deck.slides[1].shapes if shape.has_text_frame and shape.text == "Inherited title")
    assert inherited.text_frame.paragraphs[0].font.size.pt == 44
    assert inherited.text_frame.paragraphs[0].font.name == "Arial"


@pytest.mark.parametrize("name,value", [
    ("title_font_size", 0), ("title_font_size", 97), ("title_font_size", True),
    ("title_font_size", float("nan")), ("title_font_name", ""), ("title_font_name", "a" * 201),
    ("title_font_color", "red"),
])
def test_invalid_typography_is_rejected(name, value):
    with pytest.raises(ValueError):
        validate({"slides": [{"type": "content", name: value}]})


def test_typography_schema_has_groups_for_every_text_panel():
    for template in get_all_templates().values():
        assert all(field.get("group") for field in template.schema)
        for panel in template.layout:
            if panel["kind"] != "image":
                fields = {field["name"]: field for field in template.schema}
                for suffix in ("font_size", "font_name", "font_color"):
                    field = fields[panel["field"] + "_" + suffix]
                    assert field["subgroup"] == "Typography"
                    assert field["theme_key"]
def test_presentation_metadata_round_trips_and_sets_pptx_properties(client):
    metadata = {"title":"Research notes", "author":"Author", "subject":"A subject",
                "description":"A description", "keywords":"one, two", "category":"Demo", "language":"cs"}
    data = {"metadata":metadata, "slides":[{"type":"title","title":"A different slide title"}]}
    assert client.post("/api/save", json=data).status_code == 200
    assert client.get("/api/presentation").json["metadata"] == metadata
    yaml = client.post("/api/export/yaml", json=data).get_data(as_text=True)
    assert client.post("/api/import/yaml", json={"text":yaml}).json["metadata"] == metadata
    deck = Presentation(BytesIO(client.post("/api/export/pptx", json=data).data))
    for name, attribute in {"title":"title","author":"author","subject":"subject",
                            "description":"comments","keywords":"keywords","category":"category","language":"language"}.items():
        assert getattr(deck.core_properties, attribute) == metadata[name]
    assert deck.core_properties.last_modified_by == metadata["author"]
    assert deck.slides[0].shapes[0].text == "A different slide title"


@pytest.mark.parametrize("metadata", [
    None, [], "wrong", {"title":True}, {"author":123}, {"description":"x"*256}, {"title":"bad\x00title"},
])
def test_invalid_presentation_metadata_is_rejected(client, metadata):
    response = client.post("/api/import", json={"metadata":metadata,"slides":[]})
    assert response.status_code == 400


def test_custom_metadata_is_preserved():
    metadata = {"title":"Notes","custom":{"edition":3}}
    assert validate({"metadata":metadata,"slides":[]})["metadata"] == metadata
