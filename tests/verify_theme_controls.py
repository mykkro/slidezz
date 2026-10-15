"""Exercise presets, alignment, footer typography and notes online/offline."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import threading

from playwright.sync_api import sync_playwright, expect
from werkzeug.serving import make_server
from pptx import Presentation
from PIL import Image
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from backend.app import create_app, validate
from backend.services.bundles import build_static_html
from backend.services.pptx_generator import generate
from backend.templates.registry import get_all_templates

RESULTS = ROOT / "test-results"
PRESETS = json.loads((ROOT / "frontend/assets/themes.json").read_text())


def check(page, url, label):
    page.goto(url)
    expect(page.locator(".slide-item")).to_have_count(9)
    page.locator(".slide-item").nth(1).click()
    page.locator("#theme-tab").click()
    expect(page.locator("#theme-preset option")).to_have_count(6)
    page.locator("#preset-overrides").check()
    for preset in PRESETS:
        theme = preset["theme"]
        page.locator("#theme-preset").select_option(preset["id"])
        page.locator("#apply-theme-preset").click()
        page.wait_for_function("(name) => presentation.slides.every(s=>s.background_image===name)", arg=theme["default_background_image"])
        expect(page.locator("#preview .slide-canvas")).to_have_css("font-family", '"' + theme["font_name"] + '"' if " " in theme["font_name"] else theme["font_name"])
        page.wait_for_function("Array.from(document.querySelectorAll('#preview img')).every(i=>i.complete && i.naturalWidth>0)")
        assert page.evaluate("presentation.slides.every(s => !Object.hasOwn(s,'title_font_color'))")
        if label == "online": page.screenshot(path=RESULTS / ("preset-" + preset["id"] + ".png"))
    # Preserving backgrounds is independent from applying shared colors/fonts.
    old = page.evaluate("presentation.slides[0].background_image")
    page.locator("#preset-backgrounds").uncheck()
    page.locator("#theme-preset").select_option(PRESETS[0]["id"])
    page.locator("#apply-theme-preset").click()
    assert page.evaluate("presentation.slides[0].background_image") == old
    page.locator("#template-picker").click()
    page.locator("#template-menu button[data-template='image']").click()
    page.locator("#add-slide").click()
    expect(page.locator(".slide-item")).to_have_count(10)
    assert page.evaluate("current().background_image") == PRESETS[0]["theme"]["default_background_image"]
    page.locator("#theme-tab").click()
    footer = page.locator("#preview .slide-footer").first
    baseline = footer.evaluate("e=>parseFloat(e.style.top)")
    page.locator("#footer-text").fill("Shared footer")
    page.locator("#footer-font-name").select_option("PT Serif")
    page.locator("#footer-font-style").select_option("bold_italic")
    page.locator("#theme-footer_font_size").fill("20")
    page.locator("#theme-footer_margin_bottom").fill("1")
    page.locator("#footer-alignment").select_option("right")
    expect(footer).to_have_css("font-size", "26.6667px")
    expect(footer).to_have_css("font-style", "italic")
    expect(footer).to_have_css("font-weight", "700")
    expect(footer).to_have_css("text-align", "right")
    page.wait_for_function("() => parseFloat(document.querySelector('#preview .slide-footer').style.top)<600")
    assert footer.evaluate("e=>parseFloat(e.style.top)") < baseline - 50
    page.locator("#font-style").select_option("bold_italic")
    page.locator("#text-alignment").select_option("center")
    page.locator("#text-vertical-alignment").select_option("middle")
    expect(page.locator("#preview .title")).to_have_css("text-align", "center")
    expect(page.locator("#preview .title")).to_have_css("font-style", "italic")
    page.locator("#slide-tab").click()
    page.locator("#properties > details").filter(has=page.locator("#field-notes")).locator(":scope > summary").click()
    page.locator("#field-notes").fill("Private talking point, not slide content.")
    assert "Private talking point" not in page.locator("#preview").inner_text()
    title = page.locator("#properties > details").filter(has=page.locator("#field-title"))
    title.locator(".typography-group > summary").click()
    page.locator("#field-title_font_style-use-theme").uncheck()
    page.locator("#field-title_font_style button").first.click()
    expect(page.locator("#preview .title")).to_have_css("font-style", "normal")
    page.locator("#field-title_alignment-use-theme").uncheck()
    page.locator("#field-title_alignment button").nth(2).click()
    expect(page.locator("#preview .title")).to_have_css("text-align", "right")
    expect(page.locator("#save-status")).to_have_text("Saved")
    for width,height in [(1440,900),(390,844)]:
        page.set_viewport_size({"width":width,"height":height})
        assert page.evaluate("document.documentElement.scrollWidth<=innerWidth")
        page.screenshot(path=RESULTS / f"theme-controls-{label}-{width}.png")
    page.reload()
    expect(page.locator(".slide-item")).to_have_count(10)
    assert page.evaluate("presentation.slides.at(-1).notes") == "Private talking point, not slide content."
    page.locator("#editor-settings").click()
    page.locator("#editor-language").select_option("cs")
    expect(page.locator("#start-presentation")).to_have_text("Prezentovat")


def main():
    RESULTS.mkdir(exist_ok=True)
    slides = [{"type":name, **{f["name"]:f["default"] for f in t.schema if not f.get("theme_key")},
               "title_font_color":"#ff0000"} for name,t in get_all_templates().items()]
    data = validate({"slides":slides})
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        saved = root / "deck.json"
        saved.write_text(json.dumps(data), encoding="utf-8")
        offline = root / "offline.html"
        offline.write_text(build_static_html(data, {}), encoding="utf-8")
        app = create_app({"PRESENTATION_FILE":saved,"IMAGES_DIR":ROOT / "images"})
        server = make_server("127.0.0.1",0,app)
        thread = threading.Thread(target=server.serve_forever,daemon=True)
        thread.start()
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(channel="msedge")
                for label,url in [("online",f"http://127.0.0.1:{server.server_port}"),("offline",offline.as_uri())]:
                    context = browser.new_context(viewport={"width":1440,"height":900})
                    page = context.new_page()
                    errors = []
                    page.on("pageerror",lambda error: errors.append(str(error)))
                    check(page,url,label)
                    assert not errors, errors
                    context.close()
                browser.close()
        finally:
            server.shutdown(); thread.join()
        decks = []
        for preset in PRESETS:
            deck = validate({"theme":preset["theme"],"slides":[{**s,"background_image":preset["theme"]["default_background_image"]} for s in slides]})
            for slide in deck["slides"]: slide.pop("title_font_color",None)
            path = RESULTS / (preset["id"] + ".pptx")
            generate(deck,path,ROOT / "images")
            decks.append(path)
        subprocess.run([r"C:\Program Files\LibreOffice\program\soffice.exe", "-env:UserInstallation=" + (root / "office").as_uri(),
                        "--headless","--convert-to","pdf","--outdir",str(RESULTS), *map(str,decks)],check=True,timeout=180)
        sheet = Image.new("RGB", (1200, 3 * 350), "#e8eeea")
        for index,path in enumerate(decks):
            document = pymupdf.open(path.with_suffix(".pdf"))
            assert len(document) == 9
            for page in document:
                assert page.get_text().strip()
                for block in page.get_text("dict")["blocks"]:
                    if block["type"] == 0:
                        x0,y0,x1,y1 = block["bbox"]
                        assert min(x0,y0)>=0 and x1<=page.rect.width+1 and y1<=page.rect.height+1
            pix = document[1].get_pixmap()
            picture = Image.frombytes("RGB",(pix.width,pix.height),pix.samples)
            picture.thumbnail((570,321))
            sheet.paste(picture,((index%2)*600+15,(index//2)*350+15))
        sheet.save(RESULTS / "preset-pptx-contact.png")
    print("Six presets, nine layouts, notes, footer movement and typography passed online/offline and native PPTX rendering.")


if __name__ == "__main__": main()
