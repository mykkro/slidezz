from io import BytesIO
import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from backend.app import validate, create_app
from backend.services.pptx_generator import generate
from backend.services.layout_geometry import resolve_layout
from backend.templates.registry import get_template


def test_notes_and_styles_are_native_pptx_and_survive_yaml(tmp_path):
    data = validate({"theme":{"font_style":"bold_italic","text_alignment":"center","text_vertical_alignment":"middle",
                              "footer_text":"Footer","footer_font_name":"PT Serif","footer_font_size":20,
                              "footer_font_style":"bold_italic","footer_alignment":"right","footer_margin_bottom":1},
                     "slides":[{"type":"content","title":"A title","title_font_style":"regular",
                                "markdown":"Body and **emphasis**", "markdown_alignment":"right",
                                "notes":"A talking point.\nA second line."}]})
    output = BytesIO()
    generate(data, output)
    output.seek(0)
    slide = Presentation(output).slides[0]
    assert slide.notes_slide.notes_text_frame.text == data["slides"][0]["notes"]
    title, body, footer, number = [shape.text_frame for shape in slide.shapes]
    assert title.vertical_anchor == MSO_ANCHOR.MIDDLE
    assert title.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert not title.paragraphs[0].runs[0].font.bold
    assert not title.paragraphs[0].runs[0].font.italic
    assert body.paragraphs[0].alignment == PP_ALIGN.RIGHT
    assert body.paragraphs[0].runs[0].font.bold and body.paragraphs[0].runs[0].font.italic
    for frame in (footer,number):
        p = frame.paragraphs[0]
        assert p.font.name == "PT Serif" and p.font.size.pt == 20
        assert p.font.bold and p.font.italic and p.alignment == PP_ALIGN.RIGHT
    assert all("talking point" not in shape.text for shape in slide.shapes)
    client = create_app({"TESTING":True,"PRESENTATION_FILE":tmp_path / "deck.json"}).test_client()
    source = client.post("/api/export/yaml", json=data).get_data(as_text=True)
    assert client.post("/api/import/yaml",json={"text":source}).json == data


def test_footer_margin_and_size_change_geometry_without_content_overlap():
    base = validate({"slides":[]})["theme"]
    first = resolve_layout(get_template("content"), base)
    raised = resolve_layout(get_template("content"), {**base,"footer_margin_bottom":1.2,"footer_font_size":32})
    assert raised["footer"]["y"] < first["footer"]["y"] - 1
    assert raised["footer"]["h"] > first["footer"]["h"]
    assert raised["footer"]["y"] + raised["footer"]["h"] == pytest.approx(7.5-1.2)
    assert all(box["y"]+box["h"] < raised["footer"]["y"] for box in raised["layout"])


@pytest.mark.parametrize("theme", [{"footer_font_size":5},{"footer_font_size":33},{"font_style":"underline"},
                                  {"footer_font_style":False},{"text_alignment":"bad"},{"text_vertical_alignment":"center"},
                                  {"footer_alignment":None},{"footer_font_name":123}])
def test_invalid_theme_typography_rejected(theme):
    with pytest.raises(ValueError): validate({"theme":theme,"slides":[]})


@pytest.mark.parametrize("slide", [{"notes":123},{"title_alignment":"bad"},{"title_font_style":"underlined"},
                                   {"title_vertical_alignment":"center"}])
def test_invalid_panel_formatting_rejected(slide):
    with pytest.raises(ValueError): validate({"slides":[{"type":"image",**slide}]})
